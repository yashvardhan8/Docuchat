"""
Extracts text (+ metadata needed for citations) from PDF and DOCX files.

Design choice: we return a flat list of "page documents" — one entry per
PDF page (or, for DOCX, one entry for the whole doc since DOCX has no
native page concept) — each carrying enough metadata to cite later.
"""
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional

from pypdf import PdfReader
import docx as docx_lib

logger = logging.getLogger("docuchat.loader")


class EmptyDocumentError(Exception):
    pass


class CorruptedFileError(Exception):
    pass


@dataclass
class PageDocument:
    text: str
    source: str
    page: Optional[int] = None
    metadata: dict = field(default_factory=dict)


def load_pdf(path: str, display_name: str) -> List[PageDocument]:
    try:
        reader = PdfReader(path)
    except Exception as e:
        raise CorruptedFileError(f"Could not open PDF: {e}")

    pages = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            logger.warning(f"Failed to extract page {i} of {display_name}: {e}")
            text = ""
        if text.strip():
            pages.append(PageDocument(text=text, source=display_name, page=i))

    if not pages:
        raise EmptyDocumentError(
            "No extractable text found (the PDF may be scanned/image-only)."
        )
    return pages


def load_docx(path: str, display_name: str) -> List[PageDocument]:
    try:
        doc = docx_lib.Document(path)
    except Exception as e:
        raise CorruptedFileError(f"Could not open DOCX: {e}")

    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if not full_text.strip():
        raise EmptyDocumentError("The DOCX file appears to be empty.")

    # DOCX has no reliable page concept without a rendering engine, so we
    # keep page=None and rely on chunk_id for citation granularity instead.
    return [PageDocument(text=full_text, source=display_name, page=None)]


def load_document(path: str, display_name: str) -> List[PageDocument]:
    ext = Path(display_name).suffix.lower()
    if ext == ".pdf":
        return load_pdf(path, display_name)
    elif ext == ".docx":
        return load_docx(path, display_name)
    else:
        raise ValueError(f"Unsupported extension: {ext}")
